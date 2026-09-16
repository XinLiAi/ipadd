# -*- coding: utf-8 -*-
"""
GUI 无头（offscreen）集成测试
    QT_QPA_PLATFORM=offscreen python3 test_gui.py
"""

import os
import sys
import tempfile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from qt_compat import QApplication, QMessageBox, QFileDialog
from qt_compat import Qt
from qt_compat import QColor

import ip_core as C
import ip_word as W
from app_gui import MainWindow

TMP = tempfile.mkdtemp(prefix="ipgui_")

# ---- 静默弹窗 / 对话框 ----
_shown = []
QMessageBox.information = staticmethod(
    lambda *a, **k: (_shown.append(("info", str(a[1]) if len(a) > 1 else "",
                                    str(a[2]) if len(a) > 2 else "")), QMessageBox.Ok)[1])
QMessageBox.critical = staticmethod(
    lambda *a, **k: (_shown.append(("crit", str(a[1]) if len(a) > 1 else "",
                                    str(a[2]) if len(a) > 2 else "")), QMessageBox.Ok)[1])
QMessageBox.warning = staticmethod(
    lambda *a, **k: (_shown.append(("warn", str(a[1]) if len(a) > 1 else "",
                                    str(a[2]) if len(a) > 2 else "")), QMessageBox.Ok)[1])

_next_path = {"v": None}
QFileDialog.getSaveFileName = staticmethod(lambda *a, **k: (_next_path["v"], ""))
QFileDialog.getOpenFileName = staticmethod(lambda *a, **k: (_next_path["v"], ""))
QFileDialog.getExistingDirectory = staticmethod(lambda *a, **k: _next_path["v"])


def cell(tbl, r, c):
    it = tbl.item(r, c)
    return it.text() if it else None


def main():
    app = QApplication(sys.argv)
    DBPATH = os.path.join(TMP, "gui_history.db")
    w = MainWindow(db_path=DBPATH)

    print("=== 1. 窗口构造 + 内置模板自动载入 ===")
    print("   标题:", w.windowTitle())
    print("   批复单:", os.path.basename(w.ed_tpl.text()))
    print("   业务页签:", [w.cat_tabs.tabText(i) for i in range(w.cat_tabs.count())])
    assert w.ed_tpl.text().endswith(".docx"), "内置批复单模板未自动载入"
    assert w.cat_tabs.count() == 2

    print()
    print("=== 2. 自动扫描批复单 ===")
    print("   检出 %d 张纵密表" % len(w.tables))
    for t in w.tables:
        print("     %-28s %d 行" % (t.title, len(t.rows)))
    assert len(w.tables) == 4

    print()
    print("=== 3. ★自动识别（用户反馈的核心问题）===")
    print("   识别前状态:", w.lab_stat.text())
    w.do_detect()
    print("   切到业务模板:", w.cmb_tpl.currentText())
    for t in w.tables:
        print("     %-28s 起始IP=%-16s 掩码=/%d" % (
            t.title, w.row_edits[t.table_index].text(),
            w.mask_spins[t.table_index].value()))
    assert w.cmb_tpl.currentText() == "28位掩码（备用模板）"
    assert all(w.mask_spins[t.table_index].value() == 28 for t in w.tables)
    assert w.row_edits[w.tables[0].table_index].text() == "54.100.134.96"
    print("   状态栏:", w.statusBar().currentMessage())

    print()
    print("=== 4. ★非实时表是否识别到（此前一直漏掉）===")
    for t in w.tables:
        if "非实时" not in t.title:
            continue
        matched = [r for r in t.rows if r.business]
        filled = [r for r in t.rows if r.new_value]
        print("   %s" % t.title)
        print("     匹配 %d/%d 行，将填入 %d 行" % (len(matched), len(t.rows), len(filled)))
        for r in t.rows:
            flag = "✓" if r.new_value else "—"
            print("       %s %-18s 业务=%-10s 原=%-18s 新=%s" % (
                flag, r.policy, r.business or "—", r.old_value or "(空)", r.new_value or "不填"))
        # 关键断言：一/二平面非实时都必须识别到网安
        nrt_rows = {r.policy: r for r in t.rows}
        assert "区调内网管控（网安）" in nrt_rows
        assert nrt_rows["区调内网管控（网安）"].business, "非实时网安仍未识别到！"
        assert nrt_rows["区调内网管控（网安）"].new_value, "非实时网安仍未填入！"
        assert nrt_rows["区调电量"].new_value
        assert nrt_rows["区调光功率"].new_value
    print("   ✓ 一平面、二平面非实时均已识别并填入")

    print()
    print("=== 5. 逐行比对：已填的行必须与批复单原文一致 ===")
    bad = [(t.title, r.policy, r.old_value, r.new_value)
           for t in w.tables for r in t.rows
           if r.old_value and r.new_value and r.old_value != r.new_value]
    for b in bad:
        print("   ✗", b)
    assert not bad, f"{len(bad)} 行与原文不一致"
    print("   与原文不一致: 0 行 ✓")

    print()
    print("=== 6. 补填原本空着的远动地址 ===")
    new_rows = [(t.title, r.policy, r.new_value)
                for t in w.tables for r in t.rows if not r.old_value and r.new_value]
    print("   补填 %d 行" % len(new_rows))
    for x in new_rows[:3]:
        print("     ", x)
    assert len(new_rows) == 12, len(new_rows)
    assert all("远动" in p for _t, p, _v in new_rows)

    print()
    print("=== 7. 未匹配 / 未启用的业务应跳过 ===")
    skipped = [(t.title, r.policy, r.business_disabled)
               for t in w.tables for r in t.rows if not r.new_value]
    print("   跳过 %d 行" % len(skipped))
    for _t, p, d in skipped:
        print("     %-18s %s" % (p, f"（业务 {d} 未启用/无偏移）" if d else "（策略名未匹配）"))
    dis = [d for _t, _p, d in skipped if d]
    assert any("同步时钟" in d for d in dis), dis
    assert any("安控" in d for d in dis), dis

    print()
    print("=== 8. 「仅填空白处」开关 ===")
    w.chk_only_empty.setChecked(True)
    n = len([r for t in w.tables for r in t.rows if r.new_value])
    print("   仅填空白 → %d 行；状态: %s" % (n, w.lab_stat.text()))
    assert n == 12, n
    for t in w.tables:
        for r in t.rows:
            if r.old_value:
                assert not r.new_value
    w.chk_only_empty.setChecked(False)
    n2 = len([r for t in w.tables for r in t.rows if r.new_value])
    print("   关闭后 → %d 行" % n2)
    assert n2 == 44

    print()
    print("=== 9. 改起始 IP：第一接入网实时 → .160 ===")
    ti = w.tables[0].table_index
    w.row_edits[ti].setText("54.100.134.160")
    r0 = w.tables[0].rows[0]
    print("   %s: %s → %s" % (r0.policy, r0.old_value, r0.new_value))
    assert r0.new_value == "54.100.134.167-168", r0.new_value
    w.row_edits[ti].setText("54.100.134.96")

    print()
    print("=== 10. 改掩码：第一接入网实时 /28 → /27 ===")
    w.mask_spins[ti].setValue(27)
    print("   %s → %s" % (r0.policy, r0.new_value))
    assert r0.new_value == "54.100.134.103-104"   # 远动 7/8 在 /27 /28 下相同
    w.mask_spins[ti].setValue(28)

    print()
    print("=== 11. 导出 Word 批复单 ===")
    out = os.path.join(TMP, "批复单_已填.docx")
    _next_path["v"] = out
    w.do_export()
    import docx
    d = docx.Document(out)
    print("   输出 %d 字节" % os.path.getsize(out))
    # 表3=第一接入网实时，表4=第一接入网非实时，表8=第二接入网实时，表9=第二接入网非实时
    for ti_, name in ((3, "一平面实时"), (4, "一平面非实时"), (8, "二平面实时"), (9, "二平面非实时")):
        v = {}
        for ri in range(3, len(d.tables[ti_].rows)):
            c = d.tables[ti_].rows[ri].cells
            if len(c) > 4 and c[3].text.strip():
                v[c[3].text.strip()] = c[4].text.strip()
        n = len([x for x in v.values() if x])
        print("     %-12s 已填 %d/%d 行   网安=%s" % (
            name, n, len(v), v.get("区调内网管控（网安）") or v.get("区调内网管控（网安）", "—")))
        assert n > 0
    # 关键断言：非实时网安在导出文件里确实有值
    v4 = {}
    for ri in range(3, len(d.tables[4].rows)):
        c = d.tables[4].rows[ri].cells
        if len(c) > 4 and c[3].text.strip():
            v4[c[3].text.strip()] = c[4].text.strip()
    assert v4["区调内网管控（网安）"] == "54.100.135.98", v4["区调内网管控（网安）"]
    assert v4["区调电量"] == "54.100.135.106"
    print("   ✓ 一平面非实时网安 =", v4["区调内网管控（网安）"])
    v9 = {}
    for ri in range(3, len(d.tables[9].rows)):
        c = d.tables[9].rows[ri].cells
        if len(c) > 4 and c[3].text.strip():
            v9[c[3].text.strip()] = c[4].text.strip()
    assert v9["区调内网管控（网安）"] == "54.124.7.50", v9["区调内网管控（网安）"]
    print("   ✓ 二平面非实时网安 =", v9["区调内网管控（网安）"])

    print()
    print("=== 12. 批量站点导出 ===")
    w.tabs.setCurrentIndex(1)
    hdr = [w.tbl_sites.horizontalHeaderItem(c).text()
           for c in range(w.tbl_sites.columnCount())]
    print("   批量表头:", hdr)
    assert hdr[0] == "站点名称" and len(hdr) == 5
    for nm, ips in [("甲站", ["54.100.134.96", "54.100.135.96", "54.124.6.48", "54.124.7.48"]),
                    ("乙站", ["54.100.134.112", "54.100.135.112", "", ""])]:
        w._site_add()
        r = w.tbl_sites.rowCount() - 1
        w.tbl_sites.item(r, 0).setText(nm)
        for i, ip in enumerate(ips):
            w.tbl_sites.item(r, i + 1).setText(ip)
    batch_dir = os.path.join(TMP, "batch")
    os.makedirs(batch_dir, exist_ok=True)
    _next_path["v"] = batch_dir
    w.do_batch_export()
    files = sorted(os.listdir(batch_dir))
    print("   生成:", files)
    assert files == ["乙站.docx", "甲站.docx"] or set(files) == {"甲站.docx", "乙站.docx"}
    # 校验甲站用的是自己的起始IP
    da = docx.Document(os.path.join(batch_dir, "甲站.docx"))
    va = {}
    for ri in range(3, len(da.tables[3].rows)):
        c = da.tables[3].rows[ri].cells
        if len(c) > 4 and c[3].text.strip():
            va[c[3].text.strip()] = c[4].text.strip()
    db = docx.Document(os.path.join(batch_dir, "乙站.docx"))
    vb = {}
    for ri in range(3, len(db.tables[3].rows)):
        c = db.tables[3].rows[ri].cells
        if len(c) > 4 and c[3].text.strip():
            vb[c[3].text.strip()] = c[4].text.strip()
    print("   甲站 远动 =", va["备调远动fes1"])
    print("   乙站 远动 =", vb["备调远动fes1"], "(起始IP .112)")
    assert va["备调远动fes1"] == "54.100.134.103-104"
    assert vb["备调远动fes1"] == "54.100.134.119-120", vb["备调远动fes1"]
    print("   ✓ 各站点使用各自起始 IP")

    print()
    print("=== 13. 业务偏移表增删 / 停用 ===")
    w.tabs.setCurrentIndex(0)
    n0 = w.panels[0].tbl.rowCount()
    w.panels[0].add_row()
    assert w.panels[0].tbl.rowCount() == n0 + 1
    w.panels[0].tbl.selectRow(n0)
    w.panels[0].del_row()
    assert w.panels[0].tbl.rowCount() == n0
    # 停用远动 → 应显示「未启用/无偏移」且不再填入
    w.panels[0].tbl.item(0, 0).setCheckState(Qt.Unchecked)
    w.do_refresh()
    names = {r.business for t in w.tables for r in t.rows}
    dis = {r.business_disabled for t in w.tables for r in t.rows}
    print("   停用远动后：参与填充的业务里有远动吗:", "远动" in names)
    print("   标记为未启用的:", {x for x in dis if x})
    assert "远动" not in names
    assert "远动" in dis
    w.panels[0].tbl.item(0, 0).setCheckState(Qt.Checked)
    w.do_refresh()
    assert "远动" in {r.business for t in w.tables for r in t.rows}
    print("   增删/停用正常，行数回到", n0)

    print()
    print("=== 14. 配置存取 ===")
    cfg = os.path.join(TMP, "cfg.json")
    _next_path["v"] = cfg
    w.do_save_config()
    w.panels[0].ed_base.setText("172.16.9.9")
    _next_path["v"] = cfg
    w.do_load_config()
    print("   载入后实时起始IP:", w.panels[0].ed_base.text())
    assert w.panels[0].ed_base.text() != "172.16.9.9"

    print()
    print("=== 15. 非法输入容错 ===")
    w.row_edits[ti].setText("不是IP")
    w.do_refresh()
    print("   非法起始IP → 将写入行数:", w.lab_stat.text())
    w.row_edits[ti].setText("54.100.134.96")
    w.do_refresh()

    print()
    print("=== 16. 界面已精简：无独立业务地址结果表 ===")
    assert not hasattr(w, "tbl_result"), "旧的业务地址结果表应已移除"
    assert not hasattr(w, "btn_export_xlsx"), "旧的 Excel 导出按钮应已移除"
    print("   ✓ 结果表与 Excel/CSV 导出按钮已移除，界面聚焦批复单填充")

    print()
    print("=== 17. 历史台账：归档 + 查询 + 查重 ===")
    w.tabs.setCurrentIndex(0)
    w.do_detect()
    assert w.history is not None, "历史台账未初始化"

    # 顶部站点名输入框
    assert hasattr(w, "ed_site")
    w.ed_site.setText("色尼风电")
    print("   站点名称输入框:", w.ed_site.text())

    rid = w._archive("色尼风电", "/tmp/色尼风电.docx")
    print("   归档 run id =", rid)
    assert rid > 0
    st = w.history.stats()
    print("   统计:", st)
    assert st["批次数"] == 1 and st["站点数"] == 1
    assert st["地址条数"] == 28, st["地址条数"]

    # 查询
    print("   站点列表:", w.history.sites())
    assert w.history.sites() == ["色尼风电"]
    assert "远动" in w.history.businesses()
    hits = w.history.search(keyword="54.100.134.103")
    print("   按IP搜 54.100.134.103 →", len(hits), "条")
    assert len(hits) >= 1
    print("   按业务搜 远动 →", len(w.history.search(business="远动")), "条")
    print("   批次列表 →", len(w.history.list_runs()), "个")
    assert len(w.history.list_runs()) == 1

    # 界面表格是否刷新
    w._hist_reload()
    print("   批次表行数:", w.tbl_runs.rowCount(), "明细表行数:", w.tbl_hist.rowCount())
    assert w.tbl_runs.rowCount() == 1
    assert w.tbl_hist.rowCount() == 28

    # 查重：同站点应不冲突
    cf_same = w.history.find_conflicts(w.tables, exclude_site="色尼风电")
    print("   同站点重复导出 → 冲突", len(cf_same), "处（应为 0）")
    assert cf_same == []

    # 换站点、换起始 IP
    for t in w.tables:
        head = ".".join(t.suggested_base.split(".")[:3])
        w.row_edits[t.table_index].setText(head + ".112")
    w.do_refresh()
    cf_new = w.history.find_conflicts(w.tables, exclude_site="比如风电")
    print("   换到 .112 段 → 冲突", len(cf_new), "处（应为 0）")
    assert cf_new == []

    # 制造真冲突：用回原起始 IP，站点名不同
    tabs_back = w.tables
    for t in tabs_back:
        head = ".".join(t.suggested_base.split(".")[:3])
        w.row_edits[t.table_index].setText(head + ".96")
    w.do_refresh()
    cf_hit = w.history.find_conflicts(tabs_back, exclude_site="比如风电")
    print("   用回 .96 段、站点不同 → 冲突", len(cf_hit), "处")
    assert len(cf_hit) > 0
    for c in cf_hit[:3]:
        print("     %-16s %-8s 已被 %s 占用" % (c.ip, c.business, c.old_site))

    # 预览里的历史提示列
    # 先测【本站】视角：站点名仍是「色尼风电」→ 应提示「本站已记录」而非「被占用」
    w.ed_site.setText("色尼风电")
    w.do_refresh()
    w._check_history()
    notes = {}
    for tb in w.tables:
        for r in tb.rows:
            if r.hist_note:
                notes[r.hist_note] = notes.get(r.hist_note, 0) + 1
    print("   预览提示（本站视角）:", notes)
    assert any("本站已记录" in k for k in notes), notes
    assert not any("已被" in k for k in notes), notes

    # 再测【异站】视角：换成别的站点名 → 应提示「已被 色尼风电 占用」
    w.ed_site.setText("比如风电")
    w.do_refresh()
    w._check_history()
    notes2 = {}
    for tb in w.tables:
        for r in tb.rows:
            if r.hist_note:
                notes2[r.hist_note] = notes2.get(r.hist_note, 0) + 1
    print("   预览提示（异站视角）:", notes2)
    assert any("已被" in k and "色尼风电" in k for k in notes2), notes2
    print("   查重计数:", w.n_hist_conflict, "行被标记")

    # 内部重复：同业务复用不算，此处应为空
    inner = __import__("ip_ledger").self_conflicts(w.tables)
    print("   本次内部重复:", len(inner), "处（同业务复用不算，应为 0）")
    assert inner == []

    print()
    print("=== 18. 导出台账 Excel ===")
    xls = os.path.join(TMP, "台账.xlsx")
    _next_path["v"] = xls
    w._hist_export()
    print("   台账文件:", os.path.exists(xls) and "%.1f KB" % (os.path.getsize(xls) / 1024))
    assert os.path.exists(xls) and os.path.getsize(xls) > 3000

    print()
    print("=== 19. 删除 / 清空 ===")
    ids = [r["id"] for r in w.history.list_runs()]
    w.history.delete_runs(ids)
    print("   删除后批次:", len(w.history.list_runs()))
    assert w.history.list_runs() == []
    w.history.record_run(site="测试站", items=__import__("ip_ledger").collect_items(w.tables))
    w.history.clear()
    print("   清空后:", w.history.stats()["批次数"], "批次")
    assert w.history.stats()["批次数"] == 0

    print()
    print("=== 20. 数据库落盘位置 ===")
    print("   ", w.history.path)
    assert os.path.exists(w.history.path)
    assert w.history.path.endswith("history.db")

    print()
    print("=== 17. 历史台账：归档 + 查询 + 查重 ===")
    w.tabs.setCurrentIndex(0)
    w.do_detect()
    assert w.history is not None, "历史台账未初始化"

    # 顶部站点名输入框
    assert hasattr(w, "ed_site")
    w.ed_site.setText("色尼风电")
    print("   站点名称输入框:", w.ed_site.text())

    rid = w._archive("色尼风电", "/tmp/色尼风电.docx")
    print("   归档 run id =", rid)
    assert rid > 0
    st = w.history.stats()
    print("   统计:", st)
    assert st["批次数"] == 1 and st["站点数"] == 1
    assert st["地址条数"] == 28, st["地址条数"]

    # 查询
    print("   站点列表:", w.history.sites())
    assert w.history.sites() == ["色尼风电"]
    assert "远动" in w.history.businesses()
    hits = w.history.search(keyword="54.100.134.103")
    print("   按IP搜 54.100.134.103 →", len(hits), "条")
    assert len(hits) >= 1
    print("   按业务搜 远动 →", len(w.history.search(business="远动")), "条")
    print("   批次列表 →", len(w.history.list_runs()), "个")
    assert len(w.history.list_runs()) == 1

    # 界面表格是否刷新
    w._hist_reload()
    print("   批次表行数:", w.tbl_runs.rowCount(), "明细表行数:", w.tbl_hist.rowCount())
    assert w.tbl_runs.rowCount() == 1
    assert w.tbl_hist.rowCount() == 28

    # 查重：同站点应不冲突
    cf_same = w.history.find_conflicts(w.tables, exclude_site="色尼风电")
    print("   同站点重复导出 → 冲突", len(cf_same), "处（应为 0）")
    assert cf_same == []

    # 换站点、换起始 IP
    for t in w.tables:
        head = ".".join(t.suggested_base.split(".")[:3])
        w.row_edits[t.table_index].setText(head + ".112")
    w.do_refresh()
    cf_new = w.history.find_conflicts(w.tables, exclude_site="比如风电")
    print("   换到 .112 段 → 冲突", len(cf_new), "处（应为 0）")
    assert cf_new == []

    # 制造真冲突：用回原起始 IP，站点名不同
    tabs_back = w.tables
    for t in tabs_back:
        head = ".".join(t.suggested_base.split(".")[:3])
        w.row_edits[t.table_index].setText(head + ".96")
    w.do_refresh()
    cf_hit = w.history.find_conflicts(tabs_back, exclude_site="比如风电")
    print("   用回 .96 段、站点不同 → 冲突", len(cf_hit), "处")
    assert len(cf_hit) > 0
    for c in cf_hit[:3]:
        print("     %-16s %-8s 已被 %s 占用" % (c.ip, c.business, c.old_site))

    # 预览里的历史提示列
    # 先测【本站】视角：站点名仍是「色尼风电」→ 应提示「本站已记录」而非「被占用」
    w.ed_site.setText("色尼风电")
    w.do_refresh()
    w._check_history()
    notes = {}
    for tb in w.tables:
        for r in tb.rows:
            if r.hist_note:
                notes[r.hist_note] = notes.get(r.hist_note, 0) + 1
    print("   预览提示（本站视角）:", notes)
    assert any("本站已记录" in k for k in notes), notes
    assert not any("已被" in k for k in notes), notes

    # 再测【异站】视角：换成别的站点名 → 应提示「已被 色尼风电 占用」
    w.ed_site.setText("比如风电")
    w.do_refresh()
    w._check_history()
    notes2 = {}
    for tb in w.tables:
        for r in tb.rows:
            if r.hist_note:
                notes2[r.hist_note] = notes2.get(r.hist_note, 0) + 1
    print("   预览提示（异站视角）:", notes2)
    assert any("已被" in k and "色尼风电" in k for k in notes2), notes2
    print("   查重计数:", w.n_hist_conflict, "行被标记")

    # 内部重复：同业务复用不算，此处应为空
    inner = __import__("ip_ledger").self_conflicts(w.tables)
    print("   本次内部重复:", len(inner), "处（同业务复用不算，应为 0）")
    assert inner == []

    print()
    print("=== 18. 导出台账 Excel ===")
    xls = os.path.join(TMP, "台账.xlsx")
    _next_path["v"] = xls
    w._hist_export()
    print("   台账文件:", os.path.exists(xls) and "%.1f KB" % (os.path.getsize(xls) / 1024))
    assert os.path.exists(xls) and os.path.getsize(xls) > 3000

    print()
    print("=== 19. 删除 / 清空 ===")
    ids = [r["id"] for r in w.history.list_runs()]
    w.history.delete_runs(ids)
    print("   删除后批次:", len(w.history.list_runs()))
    assert w.history.list_runs() == []
    w.history.record_run(site="测试站", items=__import__("ip_ledger").collect_items(w.tables))
    w.history.clear()
    print("   清空后:", w.history.stats()["批次数"], "批次")
    assert w.history.stats()["批次数"] == 0

    print()
    print("=== 20. 数据库落盘位置 ===")
    print("   ", w.history.path)
    assert os.path.exists(w.history.path)
    assert w.history.path.endswith("history.db")

    print("\n全部 GUI 测试通过 ✓")
    print("临时目录:", TMP)


if __name__ == "__main__":
    main()
