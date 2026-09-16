# -*- coding: utf-8 -*-
"""
调度数据网接入地址批复单 —— Word 自动填充模块

作用：把工具算出的业务 IP，按「策略名称」自动填进批复单的
      「本端业务IP地址」列。

支持的批复单结构（已验证《那曲色尼风电调度数据网接入地址批复单》）：
    装置隧道/策略 表 → 表头含「策略名称」与「本端业务IP地址」两列
    一个文档可含多张此类表（如第一/第二接入网的实时、非实时纵密表）

地址格式（与批复单原文一致）：
    单网    → 54.100.134.98
    双网    → 54.100.134.103-104      （前三段相同，末位写范围）
    不连续  → 54.100.135.98、54.100.135.118
"""

from __future__ import annotations

import ipaddress
import os
import re
import shutil
from dataclasses import dataclass, field
from typing import List, Dict, Optional, Tuple

import ip_core as C

# ---------------------------------------------------------------- 策略名 → 业务 映射

# 按列表顺序匹配，策略名称（小写）包含关键词即命中，第一个命中的生效。
# 注意顺序：更具体的关键词放前面。
POLICY_RULES: List[Tuple[str, List[str]]] = [
    ("一次调频",   ["一次调频"]),
    ("一次设备",   ["一次设备"]),
    ("同步时钟",   ["同步时钟", "时钟", "对时"]),
    ("光功率",     ["光功率"]),
    ("风功率",     ["风功率"]),
    ("现货市场",   ["现货"]),
    ("网络安全监测", ["网络安全监测"]),
    ("网安",       ["内网管控", "网安"]),
    ("保信",       ["保信"]),
    ("保护",       ["保护"]),
    ("远动",       ["远动"]),
    ("PMU",        ["pmu", "同步相量"]),
    ("稳控",       ["稳控"]),
    ("电量",       ["电量", "电度"]),
    ("水情",       ["水情"]),
    ("宽频",       ["宽频"]),
    ("安控",       ["安控"]),
]

# 常用别名兜底：把业务名归一到模板里用的名称
BUSINESS_ALIAS = {
    "保信": ["保信", "保信（保护）", "保信子站"],
    "网安": ["网安", "网络安全监测"],
    "同步时钟": ["同步时钟", "时钟"],
}


def match_business(policy_name: str, available: List[str]) -> Optional[str]:
    """
    把「策略名称」映射到模板里的业务名。

    :param policy_name: 批复单里的策略名称，如 "区调PMU1"
    :param available:   当前模板里存在的业务名列表
    :return:            命中的业务名；无匹配返回 None
    """
    if not policy_name:
        return None
    low = policy_name.strip().lower()

    for biz, kws in POLICY_RULES:
        for kw in kws:
            if kw.lower() in low:
                # 归一到模板中实际存在的业务名
                for cand in BUSINESS_ALIAS.get(biz, [biz]):
                    if cand in available:
                        return cand
                if biz in available:
                    return biz
                return None       # 模板里没有该业务
    return None


# ---------------------------------------------------------------- 地址格式化

def _group_ranges(nums: List[int]) -> List[Tuple[int, int]]:
    """把 [103,104, 118] 归并为 [(103,104), (118,118)]"""
    if not nums:
        return []
    nums = sorted(set(nums))
    out = []
    start = prev = nums[0]
    for n in nums[1:]:
        if n == prev + 1:
            prev = n
        else:
            out.append((start, prev))
            start = prev = n
    out.append((start, prev))
    return out


def format_ip_list(ips: List[str]) -> str:
    """
    把业务 IP 列表格式化成批复单写法。

        ["54.100.134.98"]              → "54.100.134.98"
        ["54.100.134.103", "...104"]   → "54.100.134.103-104"
        ["54.100.135.98", "...118"]    → "54.100.135.98、54.100.135.118"
    """
    if not ips:
        return ""
    if len(ips) == 1:
        return ips[0]

    prefix3 = ".".join(ips[0].split(".")[:3])
    lasts = []
    for ip in ips:
        p = ip.split(".")
        if len(p) != 4 or ".".join(p[:3]) != prefix3:
            # 前三段不一致（跨段了），直接全部列出
            return "、".join(ips)
        try:
            lasts.append(int(p[3]))
        except ValueError:
            return "、".join(ips)

    parts = []
    for a, b in _group_ranges(lasts):
        parts.append(f"{prefix3}.{a}" if a == b else f"{prefix3}.{a}-{b}")
    return "、".join(parts)


# ---------------------------------------------------------------- 模板扫描

@dataclass
class PolicyRow:
    """批复单里一条待填的策略行"""
    row_index: int                  # 在 table.rows 中的行号
    tunnel: str = ""                # 隧道号
    policy: str = ""                # 策略名称
    old_value: str = ""             # 原有本端业务IP地址
    category: str = ""              # 所属大类（实时/非实时）
    business: Optional[str] = None  # 匹配到的业务名（None = 不会填）
    business_disabled: Optional[str] = None  # 匹配到了但该业务被停用/未分配偏移
    new_value: str = ""             # 将要填入的值
    hist_note: str = ""             # 历史占用提示（由历史记录模块填写，界面显示）
    ips: List[str] = field(default_factory=list)  # 展开后的单个 IP 列表（用于冲突检测）
    offsets: List[int] = field(default_factory=list)  # 该策略用到的偏移序号（台账用）


@dataclass
class PolicyTable:
    """批复单里一张装置隧道/策略表"""
    table_index: int                # 文档中的表格序号
    title: str                      # 表格标题（取表格前的段落文字）
    header_row: int                 # 表头所在行号
    col_policy: int                 # 「策略名称」列索引
    col_value: int                  # 「本端业务IP地址」列索引
    rows: List[PolicyRow] = field(default_factory=list)
    suggested_base: str = ""        # 从现有地址反推出的起始 IP
    suggested_mask: int = 28

    @property
    def matched(self) -> int:
        return len([r for r in rows if r.business])

    @property
    def unmatched(self) -> List[str]:
        return [r.policy for r in rows if not r.business and r.policy]


def _infer_base(values: List[str], mask: int) -> Tuple[str, int]:
    """
    从表里已有的本端业务IP反推起始 IP（网络位）。
    取出现次数最多的网络位，返回 (起始IP, 掩码)。
    """
    nets = {}
    for v in values:
        for ip in re.findall(r"\d{1,3}(?:\.\d{1,3}){3}", v or ""):
            try:
                n = ipaddress.IPv4Network(f"{ip}/{mask}", strict=False)
            except Exception:
                continue
            key = str(n.network_address)
            nets[key] = nets.get(key, 0) + 1
    if not nets:
        return "", mask
    best = max(nets.items(), key=lambda kv: kv[1])
    return best[0], mask


def scan_template(path: str) -> List[PolicyTable]:
    """
    扫描批复单，找出所有含「策略名称」+「本端业务IP地址」列的表格。
    """
    import docx

    d = docx.Document(path)
    body = d.element.body
    # 记录每个 table 前面最近的段落文字作为标题
    last_para = ""
    tables_xml = []
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            try:
                from docx.text.paragraph import Paragraph
                t = Paragraph(child, d).text.strip()
            except Exception:
                t = ""
            if t:
                last_para = t
        elif tag == "tbl":
            tables_xml.append(last_para)
            last_para = ""

    found: List[PolicyTable] = []
    for ti, tbl in enumerate(d.tables):
        if len(tbl.rows) < 2:
            continue
        # 找表头行：含「策略名称」且含「本端业务IP地址」
        hdr = col_p = col_v = None
        for ri, row in enumerate(tbl.rows):
            texts = [c.text.strip() for c in row.cells]
            if "策略名称" in texts and any("本端业务IP地址" in t for t in texts):
                hdr = ri
                col_p = texts.index("策略名称")
                col_v = next(i for i, t in enumerate(texts) if "本端业务IP地址" in t)
                break
        if hdr is None:
            continue

        rows = []
        for ri in range(hdr + 1, len(tbl.rows)):
            cells = tbl.rows[ri].cells
            if len(cells) <= max(col_p, col_v):
                continue
            policy = cells[col_p].text.strip()
            value = cells[col_v].text.strip()
            tunnel = cells[1].text.strip() if len(cells) > 1 else ""
            if not policy:
                continue     # 空行跳过
            rows.append(PolicyRow(row_index=ri, tunnel=tunnel,
                                  policy=policy, old_value=value))

        if not rows:
            continue

        base, mask = _infer_base([r.old_value for r in rows], 28)
        found.append(PolicyTable(
            table_index=ti,
            title=tables_xml[ti] if ti < len(tables_xml) else f"表{ti + 1}",
            header_row=hdr, col_policy=col_p, col_value=col_v,
            rows=rows, suggested_base=base, suggested_mask=mask,
        ))

    return found


def apply_mapping(tables: List[PolicyTable], business_names: List[str]) -> None:
    """给每张表的每行做策略名 → 业务名匹配"""
    for t in tables:
        for r in t.rows:
            r.business = match_business(r.policy, business_names)


# ---------------------------------------------------------------- 单元格写入

def _set_cell_text(cell, text: str, src_cell=None) -> None:
    """
    写入单元格文本，尽量保留原有字体格式。

    做法：保留第一个 paragraph 的第一个 run，改其文字，删掉多余 run 与段落；
          若原单元格没有 run，则从 src_cell 借一份 rPr（字体格式）。
    """
    from docx.text.paragraph import Paragraph

    paras = cell.paragraphs
    if paras:
        p = paras[0]
        runs = p.runs
        if runs:
            runs[0].text = text
            for r in runs[1:]:
                r._element.getparent().remove(r._element)
        else:
            # 没有 run：借格式
            if src_cell is not None and src_cell.paragraphs and src_cell.paragraphs[0].runs:
                proto = src_cell.paragraphs[0].runs[0]._element
                new_run = proto.__class__(proto.tag) if False else None
                import copy
                new_el = copy.deepcopy(proto)
                # 清掉原文字
                for t_el in new_el.findall(
                        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t"):
                    new_el.remove(t_el)
                from docx.oxml.ns import qn
                t_el = new_el.makeelement(qn("w:t"), {})
                t_el.text = text
                t_el.set(qn("xml:space"), "preserve")
                new_el.append(t_el)
                p._p.append(new_el)
            else:
                p.add_run(text)
        # 删除多余段落
        for extra in paras[1:]:
            extra._element.getparent().remove(extra._element)
    else:
        cell.add_paragraph(text)


def fill_docx(src_path: str, dst_path: str, tables: List[PolicyTable]) -> None:
    """
    按 tables 里已算好的 new_value 填充批复单，另存为 dst_path。
    原文件不会被修改。
    """
    import docx

    d = docx.Document(src_path)
    for t in tables:
        tbl = d.tables[t.table_index]
        for r in t.rows:
            if not r.business or not r.new_value:
                continue
            if r.row_index >= len(tbl.rows):
                continue
            cells = tbl.rows[r.row_index].cells
            if len(cells) <= t.col_value:
                continue
            # 借一个同列已有内容的单元格做格式模板
            src_cell = None
            for other in t.rows:
                if other is not r and other.old_value:
                    oc = tbl.rows[other.row_index].cells
                    if len(oc) > t.col_value:
                        src_cell = oc[t.col_value]
                        break
            _set_cell_text(cells[t.col_value], r.new_value, src_cell)

    os.makedirs(os.path.dirname(os.path.abspath(dst_path)) or ".", exist_ok=True)
    d.save(dst_path)


# ---------------------------------------------------------------- 一键计算 + 填充

def build_fill_plan(tables: List[PolicyTable],
                    category_ips: Dict[str, Dict[str, List[str]]]
                    ) -> List[str]:
    """
    为每张表算出每行应填的值。

    :param tables:       scan_template 的结果（已 apply_mapping）
    :param category_ips: {大类或表标题: {业务名: [ip, ...]}}
                         键用表标题匹配，找不到时按「实时/非实时」关键字兜底
    :return:             提示信息列表
    """
    notes = []
    for t in tables:
        pool = None
        for key, val in category_ips.items():
            if key and (key in t.title or t.title in key):
                pool = val
                break
        if pool is None:
            # 兜底：按标题里的「实时/非实时」选
            for key, val in category_ips.items():
                rt = "实时" in key and "非" not in key
                nrt = "非实时" in key
                title_rt = "实时" in t.title and "非" not in t.title
                title_nrt = "非实时" in t.title
                if (rt and title_rt) or (nrt and title_nrt):
                    pool = val
                    break
        if pool is None:
            pool = next(iter(category_ips.values()), {})

        for r in t.rows:
            if not r.business:
                r.new_value = ""
                notes.append(f"【{t.title}】策略「{r.policy}」未匹配到业务，保持原值")
                continue
            ips = pool.get(r.business)
            if not ips:
                r.new_value = ""
                notes.append(f"【{t.title}】业务「{r.business}」未分配地址，跳过")
                continue
            r.new_value = format_ip_list(ips)
    return notes



# ---------------------------------------------------------------- 大类归属

def category_for_title(title: str, categories: List[C.Category]) -> C.Category:
    """
    按表标题挑出该表该用哪个大类（实时 / 非实时）。

    标题含「非实时」→ 取名字含「非实时」的大类；否则取第一个非「非实时」的大类。
    """
    want_nrt = "非实时" in title
    for c in categories:
        if ("非实时" in c.name) == want_nrt:
            return c
    return categories[0]


def match_in_category(policy: str, cat: C.Category) -> Optional[str]:
    """
    只在该大类【已启用】的业务名里做匹配。

    两个关键点：
      1. 不跨大类 —— 避免实时「网安」与非实时「网络安全监测」撞名导致填不进去
      2. 跳过停用的业务 —— 停用即视为未分配地址，预览里显示为未匹配
    """
    return match_business(policy, [i.name for i in cat.items if i.enabled])


def apply_mapping_by_category(tables: List[PolicyTable],
                              categories: List[C.Category]) -> None:
    """
    为每张表按标题选定大类，再在该大类内做策略名匹配。

    区分两种「不会填」的情况，便于界面给出准确提示：
      - business is None 且 business_disabled is None：策略名没对上任何业务
      - business is None 且 business_disabled 有值：对上了，但该业务被停用或未分配偏移
    """
    for t in tables:
        cat = category_for_title(t.title, categories)
        t.category = cat
        enabled_names = [i.name for i in cat.items if i.enabled]
        all_names = [i.name for i in cat.items]
        for r in t.rows:
            r.category = cat.name
            r.business = match_business(r.policy, enabled_names)
            r.business_disabled = None
            if r.business is None:
                hit = match_business(r.policy, all_names)
                if hit:
                    r.business_disabled = hit


# 给 PolicyTable 动态补一个字段（避免改动上面的 dataclass 定义顺序）
PolicyTable.category = None


# ---------------------------------------------------------------- 自动识别

def _score_config(tables: List[PolicyTable], categories: List[C.Category],
                  mask: int) -> Tuple[int, int, Dict[int, str]]:
    """
    在给定掩码下试算，统计与表中现有地址吻合的行数。
    返回 (吻合数, 参与比对的行数, {table_index: 起始IP})
    """
    hit = total = 0
    bases: Dict[int, str] = {}
    for t in tables:
        cat = category_for_title(t.title, categories)
        base = infer_base([r.old_value for r in t.rows], mask)
        if not base:
            continue
        bases[t.table_index] = base
        c2 = C.Category(name=cat.name, base_network=base,
                        items=[C.BusinessItem(i.name, list(i.offsets),
                                              i.remark, i.enabled) for i in cat.items])
        res, _ = C.allocate_category(c2, mask, C.GATEWAY_NONE)
        pool: Dict[str, List[str]] = {}
        for h in res:
            pool.setdefault(h.business, []).append(h.ip)
        for r in t.rows:
            if not r.old_value:
                continue
            total += 1
            biz = match_in_category(r.policy, cat)
            ips = pool.get(biz) if biz else None
            if ips and format_ip_list(ips) == r.old_value:
                hit += 1
    return hit, total, bases


def auto_detect(tables: List[PolicyTable],
                templates: Dict[str, List[C.Category]],
                masks: Tuple[int, ...] = (25, 26, 27, 28, 29, 30)
                ) -> Optional[Dict]:
    """
    自动识别批复单用的是哪套模板 + 哪个掩码 + 各表起始 IP。

    做法：穷举 (模板 × 掩码)，用各表已有地址做吻合度打分，取最高分。
    若全部得 0（表里原本没有地址可比对），返回 None。

    :return: {"template": 模板名, "mask": 掩码, "bases": {table_index: 起始IP},
              "hit": 吻合行数, "total": 参与比对行数}
    """
    best = None
    for tname, cats in templates.items():
        for m in masks:
            hit, total, bases = _score_config(tables, cats, m)
            if total == 0:
                continue
            if best is None or hit > best["hit"]:
                best = {"template": tname, "mask": m, "bases": bases,
                        "hit": hit, "total": total}
    return best


def infer_base(values: List[str], mask: int) -> str:
    """公开版：从表里已有地址反推网络位（起始 IP）"""
    base, _m = _infer_base(values, mask)
    return base
