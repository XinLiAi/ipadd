# -*- coding: utf-8 -*-
"""
IP 地址自动分配工具 —— 导出模块
把业务地址分配结果导出为 Excel / CSV。
"""

from __future__ import annotations

import csv
import os
from typing import List

from ip_core import HostIP, Category, Issue, parse_base

# ---------------------------------------------------------------- Excel 样式

HEADER_FILL = "FFD9E1F2"      # 表头底色（淡蓝）
IP_FILL = "FFFFF2CC"          # 业务地址列底色（淡黄，突出核心列）
CONFLICT_FILL = "FFFFC7CE"    # 冲突行底色（淡红）
FONT_NAME = "微软雅黑"


def _apply_style(ws, n_col: int, n_row: int, highlight_col: int = None):
    """给 Excel 表加样式（表头加粗+底色+冻结+列宽+核心列高亮）"""
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    thin = Side(style="thin", color="FFBFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for c in range(1, n_col + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT_NAME, bold=True, size=10)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for r in range(2, n_row + 2):
        for c in range(1, n_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name=FONT_NAME, size=10)
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if highlight_col and c == highlight_col:
                cell.fill = PatternFill("solid", fgColor=IP_FILL)
                cell.font = Font(name=FONT_NAME, size=10, bold=True)

    widths = {1: 12, 2: 14, 3: 8, 4: 7, 5: 16, 6: 8, 7: 15, 8: 18,
              9: 15, 10: 24, 11: 15, 12: 10, 13: 12, 14: 16}
    for c, w in widths.items():
        if c <= n_col:
            ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = w

    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 24


def export_excel(path: str,
                 results: List[HostIP],
                 issues: List[Issue],
                 prefix: int,
                 gateway_policy: str,
                 categories: List[Category] = None) -> None:
    """
    导出分配结果到 Excel，含三个工作表：
      1. 业务地址（核心结果）
      2. 业务配置（原始偏移表，便于存档核对）
      3. 校验问题
    """
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()

    # ---------------- Sheet1 业务地址 ----------------
    ws = wb.active
    ws.title = "业务地址"

    header = HostIP.header()

    # 冲突 IP 集合，用于标红
    conflict_ips = set()
    for i in issues:
        if i.level == "error" and "都是" in i.message:
            try:
                conflict_ips.add(i.message.split("都是")[-1].strip())
            except Exception:
                pass

    ws.append(header)
    for h in results:
        ws.append(h.as_row())

    _apply_style(ws, len(header), len(results), highlight_col=5)   # 第5列=业务地址

    if conflict_ips:
        for r in range(2, len(results) + 2):
            if ws.cell(row=r, column=5).value in conflict_ips:
                for c in range(1, len(header) + 1):
                    ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=CONFLICT_FILL)

    # ---------------- Sheet2 业务配置 ----------------
    if categories:
        ws2 = wb.create_sheet("业务配置")
        ws2.append(["业务大类", "起始IP(B列基准)", "业务种类", "偏移1(C列)",
                    "偏移2(D列)", "所得值1(F列)", "所得值2(G列)", "启用", "备注"])
        for cat in categories:
            base = parse_base(cat.base_network)
            b = base[3] if base else 0
            for it in cat.items:
                o1 = it.offsets[0] if len(it.offsets) > 0 else ""
                o2 = it.offsets[1] if len(it.offsets) > 1 else ""
                f1 = (b + o1) if o1 != "" else ""
                f2 = (b + o2) if o2 != "" else ""
                ws2.append([cat.name, cat.base_network, it.name, o1, o2, f1, f2,
                            "是" if it.enabled else "否", it.remark])
        _apply_style(ws2, 9, sum(len(c.items) for c in categories), highlight_col=6)

    # ---------------- Sheet3 校验问题 ----------------
    ws3 = wb.create_sheet("校验问题")
    ws3.append(["级别", "说明"])
    if issues:
        for i in issues:
            ws3.append(["错误" if i.level == "error" else "警告", i.message])
    else:
        ws3.append(["通过", "未发现地址冲突或配置错误"])
    _apply_style(ws3, 2, max(1, len(issues)))

    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
    wb.save(path)


def export_csv(path: str, results: List[HostIP]) -> None:
    """导出为 CSV（UTF-8 BOM，Excel 直接双击不乱码）"""
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(HostIP.header())
        for h in results:
            w.writerow(h.as_row())


def export_dict_rows_excel(path: str, rows: List[dict], headers: List[str],
                           sheet_title: str = "批量业务地址") -> None:
    """把批量站点结果（dict 列表）导出为 Excel"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h, "") for h in headers])
    _apply_style(ws, len(headers), len(rows),
                 highlight_col=(headers.index("业务地址") + 1) if "业务地址" in headers else None)
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
    wb.save(path)


def export_dict_rows_csv(path: str, rows: List[dict], headers: List[str]) -> None:
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(headers)
        for r in rows:
            w.writerow([r.get(h, "") for h in headers])
